"""The Excel workbook: a live-formula dashboard over ``tblApplications`` + SQL-fed report sheets.

Two kinds of numbers live in the workbook, and every sheet says which it holds:

* **Live formulas** (Dashboard): COUNTIFS / SUMPRODUCT / MEDIAN / INDEX-MATCH over the
  ``tblApplications`` Excel Table and the ``Benchmarks`` / ``ResponseBands`` tables. They
  recalculate when rows are pasted in, and ``tests/test_bi.py`` recalculates them with
  LibreOffice and asserts they equal the SQL metric layer.
* **SQL snapshots** (Funnel, Platforms, Weekly, ...): the output of
  ``jobtracker.warehouse.metrics`` written as values - window functions, Kaplan-Meier and
  Wilson intervals are computed once, in SQL, never re-derived here.

Only Excel-2007-era functions are used (no LET / FILTER / XLOOKUP / dynamic arrays), so the
file calculates identically in Excel 2010+, Excel 365 and LibreOffice.
"""
from __future__ import annotations

import io
import math
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
)
from openpyxl.drawing.text import Font as DrawingFont
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from .. import benchmarks as bm
from ..warehouse.metrics import SQL, Metrics, Z, rate, wilson
from . import catalog, style
from .dax import measures as dax_measure_list
from .tableau import calculations as tableau_calc_list

# ----------------------------------------------------------------------------- styling

FONT = style.BODY_FONT
HEAD = style.HEADING_FONT


def _font(size=10, bold=False, color=style.INK, italic=False, name: str = FONT) -> Font:
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def _fill(hex_colour: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_colour, end_color=hex_colour)


# Light editorial look: paper / ink, hairline rules, one accent used sparingly.
THIN = Side(style="thin", color=style.RULE)
INK_RULE = Side(style="thin", color=style.INK)
ACCENT_RULE = Side(style="medium", color=style.ACCENT)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = _font(10, True, style.INK)
HEADER_FILL = _fill(style.PAPER)
HEADER_BORDER = Border(top=INK_RULE, bottom=INK_RULE)
TITLE_FONT = _font(20, True, style.INK, name=HEAD)
SUBTITLE_FONT = _font(10, False, style.SECONDARY)
PAPER_FILL = _fill(style.PAPER)
SURFACE_FILL = _fill(style.SURFACE)
SECTION_FONT = _font(13, True, style.INK, name=HEAD)
NOTE_FONT = _font(9, False, style.MUTED, italic=True)
LINK_FONT = Font(name=FONT, size=10, color=style.ACCENT, underline="single")
INPUT_FONT = _font(10, True, style.ACCENT)
INPUT_FILL = _fill(style.ACCENT_SOFT)
TABLE_STYLE = "TableStyleLight1"

FMT_INT = "#,##0"
FMT_PCT = "0.0%"
FMT_PCT_NUM = "0.0"
FMT_DATE = "yyyy-mm-dd"
FMT_DAYS = "0.0"

SHEETS = ["Cover", "Dashboard", "Applications", "Funnel", "Platforms", "Weekly", "Cohorts",
          "Reply time", "Follow-ups", "Segments", "Data quality", "Data dictionary", "Benchmarks"]

# Applications sheet: the most useful columns first; any new rpt column is appended.
APP_COLUMN_ORDER = [
    "applied_date", "company_name", "role_title", "platform", "channel", "current_status",
    "funnel_stage", "max_stage_reached", "role_family", "seniority", "industry", "job_type",
    "location", "is_response", "reached_assessment", "reached_interview", "reached_offer",
    "is_rejected", "is_ghosted", "is_active", "is_terminal", "is_direct", "first_response_date",
    "last_activity_date", "days_to_response", "days_to_interview", "days_since_applied",
    "days_since_activity", "applied_week", "applied_month", "applied_year", "applied_quarter",
    "applied_weekday", "applied_weekday_num", "applied_on_weekend", "event_count",
    "has_job_description", "has_low_trust_dates", "origin", "parsed_by", "parse_confidence",
    "job_url", "gmail_link", "application_id", "application_count",
]


def _py(value):
    """pandas / numpy scalar -> plain Python (None for every flavour of missing)."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.date() if value == value.normalize() else value.to_pydatetime()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def _as_text(cell, text: str) -> None:
    """Write text that may start with '=' as a literal string, not a formula."""
    cell.value = text
    cell.data_type = "s"


def _rich(size: int, colour: str, bold: bool = False) -> RichText:
    cp = CharacterProperties(sz=size * 100, b=bold, solidFill=colour,
                             latin=DrawingFont(typeface=FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title_text(title) -> str:
    if title is None:
        return ""
    if isinstance(title, str):
        return title
    try:
        return "".join(run.t for para in title.tx.rich.p for run in (para.r or []))
    except AttributeError:
        return ""


# ----------------------------------------------------------------------------- data

@dataclass
class _Data:
    as_of: str
    built_at: str
    frames: dict[str, pd.DataFrame]
    kpis: dict
    applications: pd.DataFrame
    follow_up_days: int
    ghost_after_days: int


def _collect(con, settings) -> _Data:
    metrics = Metrics(con, settings)
    names = ["funnel", "platform_performance", "weekly_velocity", "monthly_cohorts",
             "response_survival", "response_time_buckets", "stage_timing", "follow_up_queue",
             "data_quality", "status_mix", "period_compare"]
    frames = {n: metrics.query(n) for n in names}
    for seg in ("role_family", "industry", "seniority", "job_type"):
        frames[f"segment_{seg}"] = metrics.query("segment_performance", segment=seg)
    apps = con.execute(
        f"SELECT * FROM {catalog.FLAT_VIEW} ORDER BY applied_date DESC, company_name, application_id"
    ).df()
    try:
        import json
        row = con.execute("SELECT manifest FROM _build_manifest").fetchone()
        built_at = json.loads(row[0]).get("built_at", "") if row else ""
    except Exception:
        built_at = ""
    built_at = built_at or datetime.now().isoformat(timespec="seconds")
    return _Data(
        as_of=built_at[:10], built_at=built_at.replace("T", " "), frames=frames,
        kpis=metrics.kpis(), applications=apps,
        follow_up_days=int(getattr(settings, "follow_up_after_days", 7)),
        ghost_after_days=int(getattr(settings, "ghost_after_days", 21)),
    )


# ----------------------------------------------------------------------------- KPI registry

@dataclass(frozen=True)
class Kpi:
    key: str            # defined name suffix: kpi_<key>
    label: str
    formula: str
    fmt: str
    sql_column: str     # column of the SQL `kpis` / `period_compare` query it must equal
    definition: str = ""
    sql: str = ""
    dax: str = ""       # measure name
    tableau: str = ""   # calculated field name
    scale: float = 1.0  # multiply the Excel value by this to compare with SQL (fractions -> pct)


T = "tblApplications"
_RESP = "COUNT(*) FILTER (WHERE is_response)"
_INT = "COUNT(*) FILTER (WHERE reached_interview)"


def _wilson_xl(rate_name: str, bound: str) -> str:
    sign = "-" if bound == "lo" else "+"
    n = "kpi_applications"
    p = rate_name
    return (f"=IF({n}>0,MAX(0,MIN(1,({p}+WilsonZ^2/(2*{n}){sign}WilsonZ*SQRT({p}*(1-{p})/{n}"
            f"+WilsonZ^2/(4*{n}^2)))/(1+WilsonZ^2/{n}))),\"\")")


def _count_true(col: str) -> str:
    return f"=COUNTIFS({T}[{col}],TRUE)"


def _rate_xl(numerator: str) -> str:
    return f'=IF(kpi_applications>0,{numerator}/kpi_applications,"")'


KPIS: list[Kpi] = [
    Kpi("applications", "Applications", f"=COUNTA({T}[application_id])", FMT_INT, "applications",
        "Applications sent (one row of tblApplications each).", "COUNT(*)",
        "Total Applications", "Total Applications"),
    Kpi("responses", "Employer responses", _count_true("is_response"), FMT_INT, "responses",
        "Applications with any employer reply: assessment, interview, offer or rejection.",
        _RESP, "Responses", "Responses"),
    Kpi("assessments", "Reached assessment", _count_true("reached_assessment"), FMT_INT,
        "assessments", "Reached an assessment or further (furthest stage).",
        "COUNT(*) FILTER (WHERE reached_assessment)", "Assessments", "Assessments"),
    Kpi("interviews", "Reached interview", _count_true("reached_interview"), FMT_INT, "interviews",
        "Reached an interview or further - a later rejection does not erase it.", _INT,
        "Interviews", "Interviews"),
    Kpi("offers", "Offers", _count_true("reached_offer"), FMT_INT, "offers",
        "Offers received.", "COUNT(*) FILTER (WHERE reached_offer)", "Offers", "Offers"),
    Kpi("rejections", "Rejected", _count_true("is_rejected"), FMT_INT, "rejections",
        "Current status Rejected.", "COUNT(*) FILTER (WHERE is_rejected)", "Rejections",
        "Rejections"),
    Kpi("ghosted", "Ghosted", _count_true("is_ghosted"), FMT_INT, "ghosted",
        "Current status Ghosted (no reply within the ghost window).",
        "COUNT(*) FILTER (WHERE is_ghosted)", "Ghosted", "Ghosted Applications"),
    Kpi("active", "Active pipeline", _count_true("is_active"), FMT_INT, "active",
        "Still open: Applied, Assessment / OA or Interview.", "COUNT(*) FILTER (WHERE is_active)",
        "Active Pipeline", "Active Pipeline"),
    Kpi("needs_follow_up", "Needs follow-up",
        f'=COUNTIFS({T}[is_active],TRUE,{T}[days_since_activity],">="&FollowUpDays)', FMT_INT,
        "needs_follow_up", "Open applications silent for at least FollowUpDays days.",
        "COUNT(*) FILTER (WHERE is_active AND days_since_activity >= $follow_up_days)",
        "Needs Follow-up", "Needs Follow-up"),
    Kpi("companies", "Companies",
        f"=IFERROR(SUMPRODUCT(({T}[company_name]<>\"\")/COUNTIF({T}[company_name],{T}[company_name]&\"\")),0)",
        FMT_INT, "companies", "Distinct employers (SUMPRODUCT / COUNTIF distinct count).",
        "COUNT(DISTINCT company_name)", "Companies", "Companies"),
    Kpi("response_rate", "Response rate", _rate_xl("kpi_responses"), FMT_PCT, "response_rate_pct",
        "Employer responses / applications.", rate(_RESP, "COUNT(*)"), "Response Rate %",
        "Response Rate %", 100),
    Kpi("response_rate_lo", "Response rate - 95% CI low", _wilson_xl("kpi_response_rate", "lo"),
        FMT_PCT, "response_rate_lo", "Wilson score interval lower bound (z = WilsonZ).",
        wilson(_RESP, "COUNT(*)", "lo"), "Response Rate CI Low %", "Response Rate CI Low %", 100),
    Kpi("response_rate_hi", "Response rate - 95% CI high", _wilson_xl("kpi_response_rate", "hi"),
        FMT_PCT, "response_rate_hi", "Wilson score interval upper bound.",
        wilson(_RESP, "COUNT(*)", "hi"), "Response Rate CI High %", "Response Rate CI High %", 100),
    Kpi("interview_rate", "Interview rate", _rate_xl("kpi_interviews"), FMT_PCT,
        "interview_rate_pct", "Reached interview / applications.", rate(_INT, "COUNT(*)"),
        "Interview Rate %", "Interview Rate %", 100),
    Kpi("interview_rate_lo", "Interview rate - 95% CI low", _wilson_xl("kpi_interview_rate", "lo"),
        FMT_PCT, "interview_rate_lo", "Wilson lower bound of the interview rate.",
        wilson(_INT, "COUNT(*)", "lo"), "Interview Rate CI Low %", "Interview Rate CI Low %", 100),
    Kpi("interview_rate_hi", "Interview rate - 95% CI high", _wilson_xl("kpi_interview_rate", "hi"),
        FMT_PCT, "interview_rate_hi", "Wilson upper bound of the interview rate.",
        wilson(_INT, "COUNT(*)", "hi"), "Interview Rate CI High %", "Interview Rate CI High %", 100),
    Kpi("offer_rate", "Offer rate", _rate_xl("kpi_offers"), FMT_PCT, "offer_rate_pct",
        "Offers / applications.", rate("COUNT(*) FILTER (WHERE reached_offer)", "COUNT(*)"),
        "Offer Rate %", "Offer Rate %", 100),
    Kpi("rejection_rate", "Rejection rate", _rate_xl("kpi_rejections"), FMT_PCT,
        "rejection_rate_pct", "Rejected / applications.",
        rate("COUNT(*) FILTER (WHERE is_rejected)", "COUNT(*)"), "Rejection Rate %",
        "Rejection Rate %", 100),
    Kpi("ghost_rate", "Ghost rate", _rate_xl("kpi_ghosted"), FMT_PCT, "ghost_rate_pct",
        "Ghosted / applications.", rate("COUNT(*) FILTER (WHERE is_ghosted)", "COUNT(*)"),
        "Ghost Rate %", "Ghost Rate %", 100),
    Kpi("median_days_to_response", "Median days to reply",
        f'=IF(COUNT({T}[days_to_response])>0,MEDIAN({T}[days_to_response]),"")', FMT_DAYS,
        "median_days_to_response",
        "Median over KNOWN reply dates: MEDIAN skips the blank cells, and unknown dates are "
        "left blank, never 0.", "MEDIAN(days_to_response)", "Median Days To Response",
        "Median Days To Response"),
    Kpi("avg_days_to_response", "Average days to reply",
        f'=IFERROR(AVERAGEIFS({T}[days_to_response],{T}[days_to_response],"<>"),"")', FMT_DAYS,
        "avg_days_to_response", "Mean over known reply dates (AVERAGEIFS on non-blank cells).",
        "ROUND(AVG(days_to_response), 1)", "Average Days To Response", "Avg Days To Response"),
    Kpi("known_reply_dates", "Replies with a known date", f"=COUNT({T}[days_to_response])",
        FMT_INT, "responses_with_known_date", "COUNT ignores blanks: replies whose date is known.",
        "COUNT(days_to_response)", "Responses With Known Reply Date",
        "Responses With Known Reply Date"),
    Kpi("median_days_to_interview", "Median days to interview",
        f'=IF(COUNT({T}[days_to_interview])>0,MEDIAN({T}[days_to_interview]),"")', FMT_DAYS,
        "median_days_to_interview", "Median days to first interview (known dates only).",
        "MEDIAN(days_to_interview)", "Median Days To Interview", "Median Days To Interview"),
    Kpi("direct_share", "Direct channel share",
        _rate_xl(f"COUNTIFS({T}[is_direct],TRUE)"), FMT_PCT, "direct_share_pct",
        "Share applied directly with the employer (careers site / ATS / referral).",
        rate("COUNT(*) FILTER (WHERE is_direct)", "COUNT(*)"), "Direct Channel Share %",
        "Direct Channel Share %", 100),
    Kpi("response_band", "Response band",
        '=IF(kpi_applications>0,INDEX(ResponseBands[band],'
        'MATCH(ROUND(kpi_response_rate*100,1),ResponseBands[lower_pct],1)),"No data")', "General",
        "response_band",
        "Band of the response rate (rounded to 1 dp) from the ResponseBands table: approximate "
        "MATCH finds the highest lower bound <= the rate.",
        bm.response_band_sql(rate(_RESP, "COUNT(*)")), "Response Band", "Response Band"),
    Kpi("interview_vs_direct_pp", "Interview rate vs direct benchmark (pp)",
        '=IF(kpi_applications>0,ROUND(kpi_interview_rate*100-DirectInterviewBenchmark,1),"")',
        "+0.0;-0.0;0.0", "",
        f"Interview rate minus the {bm.DIRECT_INTERVIEW_RATE_PCT}% direct-application benchmark.",
        f"ROUND(100.0 * {_INT} / COUNT(*) - {bm.DIRECT_INTERVIEW_RATE_PCT}, 1)",
        "Interview Rate vs Direct Benchmark (pp)", "Interview Rate vs Direct Benchmark (pp)"),
    Kpi("sample_confidence", "Sample confidence",
        '=IF(kpi_applications>=SampleAdequate,"Adequate",'
        'IF(kpi_applications>=SampleIndicative,"Indicative","Too small"))', "General", "",
        "How much weight the sample can bear.",
        f"CASE WHEN COUNT(*) >= {bm.SAMPLE_ADEQUATE} THEN 'Adequate' WHEN COUNT(*) >= "
        f"{bm.SAMPLE_INDICATIVE} THEN 'Indicative' ELSE 'Too small' END",
        "Sample Confidence", "Sample Confidence"),
    Kpi("last_30", "Applications - last 30 days", f'=COUNTIFS({T}[days_since_applied],"<30")',
        FMT_INT, "last_30", "Applied in the 30 days up to the build date.",
        "COUNT(*) FILTER (WHERE applied_date > current_date - 30)", "Applications Last 30 Days",
        "Applications Last 30 Days"),
    Kpi("prev_30", "Applications - previous 30 days",
        f'=COUNTIFS({T}[days_since_applied],">=30",{T}[days_since_applied],"<60")', FMT_INT,
        "prev_30", "Applied 30-59 days before the build date.",
        "COUNT(*) FILTER (WHERE applied_date > current_date - 60 AND applied_date <= current_date - 30)",
        "Applications Prev 30 Days", "Applications Prev 30 Days"),
    Kpi("change_30", "Change vs previous 30 days", "=kpi_last_30-kpi_prev_30", "+#,##0;-#,##0;0",
        "change", "Last 30 days minus the previous 30.", "last_30 - prev_30",
        "Applications Change 30d", "Applications Change 30d"),
    Kpi("first_applied", "First application", f"=MIN({T}[applied_date])", FMT_DATE,
        "first_applied", "Earliest applied_date.", "MIN(applied_date)"),
    Kpi("last_applied", "Latest application", f"=MAX({T}[applied_date])", FMT_DATE,
        "last_applied", "Latest applied_date.", "MAX(applied_date)"),
]

# Benchmarks table rows: (key, defined name, label, value, unit, source)
def _benchmark_rows(follow_up_days: int, ghost_after_days: int) -> list[tuple]:
    r_lo, r_hi = bm.FUNNEL_PER_100["Employer responded"]
    i_lo, i_hi = bm.FUNNEL_PER_100["Interview"]
    o_lo, o_hi = bm.FUNNEL_PER_100["Offer"]
    src = "jobtracker/benchmarks.py"
    return [
        ("direct_interview_rate_pct", "DirectInterviewBenchmark",
         "Interview rate - direct applications", bm.DIRECT_INTERVIEW_RATE_PCT, "%",
         f"{src}: 1.24M tracked applications"),
        ("job_board_interview_rate_pct", "JobBoardInterviewBenchmark",
         "Interview rate - job-board quick apply", bm.JOB_BOARD_INTERVIEW_RATE_PCT, "%", src),
        ("responses_per_100_lo", "BenchRespLo", "Employer responses per 100 - low", r_lo,
         "per 100 applications", f"{src}: published 100-application funnel"),
        ("responses_per_100_hi", "BenchRespHi", "Employer responses per 100 - high", r_hi,
         "per 100 applications", src),
        ("interviews_per_100_lo", "BenchIntLo", "Interviews per 100 - low", i_lo,
         "per 100 applications", src),
        ("interviews_per_100_hi", "BenchIntHi", "Interviews per 100 - high", i_hi,
         "per 100 applications", src),
        ("offers_per_100_lo", "BenchOfferLo", "Offers per 100 - low", o_lo, "per 100 applications",
         src),
        ("offers_per_100_hi", "BenchOfferHi", "Offers per 100 - high", o_hi, "per 100 applications",
         src),
        ("sample_adequate", "SampleAdequate", "Sample size for 'Adequate'", bm.SAMPLE_ADEQUATE,
         "applications", src),
        ("sample_indicative", "SampleIndicative", "Sample size for 'Indicative'",
         bm.SAMPLE_INDICATIVE, "applications", src),
        ("wilson_z", "WilsonZ", "z for the 95% Wilson interval", Z, "z",
         "jobtracker/warehouse/metrics.py (Z)"),
        ("follow_up_days", "FollowUpDays", "Follow up after (days of silence)", follow_up_days,
         "days", "settings.follow_up_after_days - INPUT, edit freely"),
        ("ghost_after_days", "GhostAfterDays", "Mark ghosted after (days)", ghost_after_days,
         "days", "settings.ghost_after_days (used by the sync, shown for reference)"),
    ]


# ----------------------------------------------------------------------------- builder

class _Builder:
    def __init__(self, data: _Data):
        self.d = data
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets = {name: self.wb.create_sheet(name) for name in SHEETS}
        self.app_cols: list[str] = []
        self.n_apps = len(data.applications)
        self.platform_rows: list[str] = []

    # ------------------------------------------------------------------ helpers
    def _name(self, name: str, sheet: str, ref: str) -> None:
        self.wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!{ref}")

    def _banner(self, ws, title: str, subtitle: str, last_col: int = 12) -> None:
        ws.sheet_view.showGridLines = False
        for col in range(1, last_col + 1):
            ws.cell(row=2, column=col).border = Border(bottom=INK_RULE)
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 18
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws["A1"].alignment = Alignment(vertical="center")
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE_FONT

    def _print(self, ws, landscape=True, title_rows: str | None = None) -> None:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        ws.oddHeader.left.text = "&A"
        ws.oddHeader.right.text = f"Job search analytics - as of {self.d.as_of}"
        ws.oddFooter.center.text = "Page &P of &N"
        if title_rows:
            ws.print_title_rows = title_rows

    def _write_table(self, ws, frame: pd.DataFrame, top: int, left: int, name: str,
                     formats: dict[str, str] | None = None, widths: Mapping[str, float] | None = None,
                     style_name: str = TABLE_STYLE) -> tuple[int, int, int, int]:
        """Write ``frame`` as an Excel Table; returns (first_row, last_row, first_col, last_col)."""
        formats = formats or {}
        cols = list(frame.columns)
        for j, col in enumerate(cols):
            cell = ws.cell(row=top, column=left + j, value=str(col))
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = HEADER_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        rows = frame.itertuples(index=False, name=None)
        n = 0
        for i, values in enumerate(rows, start=1):
            n = i
            for j, value in enumerate(values):
                value = _py(value)
                cell = ws.cell(row=top + i, column=left + j, value=value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"  # data, never a formula (formula-injection guard)
                cell.font = _font()
                fmt = formats.get(cols[j])
                if fmt:
                    cell.number_format = fmt
                elif isinstance(cell.value, date):
                    cell.number_format = FMT_DATE
        if n == 0:  # an Excel Table needs one data row
            n = 1
        last_row = top + n
        last_col = left + len(cols) - 1
        ref = f"{get_column_letter(left)}{top}:{get_column_letter(last_col)}{last_row}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name=style_name, showRowStripes=True,
                                              showColumnStripes=False)
        ws.add_table(table)
        for j, col in enumerate(cols):
            letter = get_column_letter(left + j)
            width = (widths or {}).get(col)
            if width is None:
                sample = [len(str(_py(v))) for v in frame[col].head(200) if _py(v) is not None]
                width = min(42, max(len(str(col)) + 4, (max(sample) + 2) if sample else 8))
            current = ws.column_dimensions[letter].width or 0
            ws.column_dimensions[letter].width = max(current, width)
        return top + 1, last_row, left, last_col

    def _section(self, ws, cell: str, title: str, note: str = "") -> None:
        ws[cell] = title
        ws[cell].font = SECTION_FONT
        if note:
            below = ws[cell].offset(row=1)
            below.value = note
            below.font = NOTE_FONT

    # ------------------------------------------------------------------ sheets
    def build(self) -> Workbook:
        self._benchmarks()
        self._applications()
        self._dashboard()
        self._funnel()
        self._platforms()
        self._weekly()
        self._cohorts()
        self._reply_time()
        self._follow_ups()
        self._segments()
        self._data_quality()
        self._dictionary()
        self._cover()
        self.wb.active = 0
        props = self.wb.properties
        props.title = "Job search analytics workbook"
        props.subject = "Job-application pipeline: star schema, KPIs, benchmarks"
        props.creator = "jobtracker"
        props.keywords = "SQL, Excel, Power BI, Tableau, star schema"
        return self.wb

    # ---- Benchmarks ---------------------------------------------------------------
    def _benchmarks(self) -> None:
        ws = self.sheets["Benchmarks"]
        self._banner(ws, "Benchmarks & settings",
                     "Generated from jobtracker/benchmarks.py - the same constants the SQL, DAX "
                     "and Tableau layers use. Blue cells on a blue tint are inputs.", 9)
        rows = _benchmark_rows(self.d.follow_up_days, self.d.ghost_after_days)
        frame = pd.DataFrame([r[:1] + r[2:] for r in rows],
                             columns=["key", "benchmark", "value", "unit", "source"])
        first, last, _, _ = self._write_table(
            ws, frame, 4, 1, "Benchmarks",
            widths={"key": 30, "benchmark": 40, "value": 10, "unit": 20, "source": 52})
        for i, row in enumerate(rows):
            r = first + i
            self._name(row[1], "Benchmarks", f"$C${r}")
            ws.cell(row=r, column=3).number_format = "0.00" if isinstance(row[3], float) else "0"
            if row[0] == "follow_up_days":
                ws.cell(row=r, column=3).font = INPUT_FONT
                ws.cell(row=r, column=3).fill = INPUT_FILL
        bands = sorted(bm.RESPONSE_RATE_BANDS, key=lambda b: b.lower_pct)
        labels = []
        for i, band in enumerate(bands):
            upper = bands[i + 1].lower_pct if i + 1 < len(bands) else None
            labels.append(f"< {upper:g}%" if band.lower_pct == 0 and upper is not None
                          else f"{band.lower_pct:g}%+" if upper is None
                          else f"{band.lower_pct:g}-{upper:g}%")
        band_frame = pd.DataFrame({"band": [b.label for b in bands],
                                   "lower_pct": [b.lower_pct for b in bands],
                                   "range": labels})
        self._write_table(ws, band_frame, 4, 7, "ResponseBands",
                          formats={"lower_pct": "0.0"},
                          widths={"band": 16, "lower_pct": 11, "range": 10})
        ws["G3"] = "Response-rate bands (ascending - approximate MATCH needs this order)"
        ws["G3"].font = NOTE_FONT
        ws["A3"] = ("Named cells: " + ", ".join(r[1] for r in rows)
                    + ". Editing a value here changes this workbook only; the app, DAX and "
                      "Tableau read jobtracker/benchmarks.py.")
        ws["A3"].font = NOTE_FONT
        note_row = last + 2
        ws.cell(row=note_row, column=1, value=(
            "Caveat: these are aggregate figures from public job-search research (mixed seniority, "
            "mostly cold applications). One person's pipeline is a small sample - comparisons are "
            "indicative, which is why every rate carries a Wilson 95% interval.")).font = NOTE_FONT
        ws.freeze_panes = "A5"
        self._print(ws)

    # ---- Applications ---------------------------------------------------------------
    def _applications(self) -> None:
        ws = self.sheets["Applications"]
        frame = self.d.applications
        keys = {c for c in frame.columns if c.endswith("_key")}
        order = [c for c in APP_COLUMN_ORDER if c in frame.columns]
        order += [c for c in frame.columns if c not in order and c not in keys]
        frame = frame[order]
        self.app_cols = order
        formats = {}
        widths = {}
        for col in order:
            kind = catalog.column_type(col)
            if kind == "date":
                formats[col] = FMT_DATE
                widths[col] = 12
            elif kind == "boolean":
                widths[col] = max(9, len(col) + 3)
            elif kind == "integer":
                formats[col] = "0"
                widths[col] = max(9, len(col) + 3)
            elif kind == "decimal":
                formats[col] = "0.00"
                widths[col] = max(9, len(col) + 3)
        widths.update({"company_name": 28, "role_title": 34, "job_url": 40, "gmail_link": 30,
                       "application_id": 22, "location": 16, "current_status": 17})
        first, last, _, _ = self._write_table(ws, frame, 1, 1, "tblApplications",
                                                     formats, widths)
        ws.freeze_panes = "D2"
        ws.row_dimensions[1].height = 30
        letter = {c: get_column_letter(i + 1) for i, c in enumerate(order)}

        # hyperlinks for the link columns
        for col in ("job_url", "gmail_link"):
            if col not in letter:
                continue
            for r in range(first, last + 1):
                cell = ws[f"{letter[col]}{r}"]
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = LINK_FONT
        # status chips
        if "current_status" in letter:
            rng = f"{letter['current_status']}{first}:{letter['current_status']}{last}"
            for status, (bg, fg) in style.STATUS_CHIPS.items():
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=[f'"{status}"'], fill=_fill(bg),
                    font=Font(color=fg, bold=True)))
        # silence heat: 0 -> follow-up threshold -> ghost window
        if "days_since_activity" in letter:
            rng = f"{letter['days_since_activity']}{first}:{letter['days_since_activity']}{last}"
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="num", start_value=0, start_color=style.SURFACE,
                mid_type="num", mid_value=self.d.follow_up_days,
                mid_color=style.tint(style.CAUTION, 0.75),
                end_type="num", end_value=self.d.ghost_after_days,
                end_color=style.tint(style.NEGATIVE, 0.65)))
        for col, colour in (("days_to_response", style.STATUS["Assessment / OA"]),
                            ("days_since_applied", style.CONTEXT)):
            if col in letter:
                rng = f"{letter[col]}{first}:{letter[col]}{last}"
                ws.conditional_formatting.add(rng, DataBarRule(
                    start_type="num", start_value=0, end_type="max", color=colour,
                    showValue=True))
        # reached_* flags: green when TRUE
        for col in ("is_response", "reached_assessment", "reached_interview", "reached_offer"):
            if col in letter:
                rng = f"{letter[col]}{first}:{letter[col]}{last}"
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=["TRUE"], font=Font(color=style.POSITIVE, bold=True)))
        ws.sheet_view.zoomScale = 90
        self._print(ws, title_rows="1:1")

    # ---- Dashboard ------------------------------------------------------------------
    def _dashboard(self) -> None:
        ws = self.sheets["Dashboard"]
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.column_dimensions["A"].width = 2
        for letter in "BCDEFGHI":
            ws.column_dimensions[letter].width = 14
        ws.column_dimensions["J"].width = 3
        for letter, width in zip("KLMNOPQ", (34, 14, 30, 15, 15, 15, 13), strict=True):
            ws.column_dimensions[letter].width = width
        # paper canvas; tiles, tables and charts sit on it as white surfaces
        for row in range(1, 60):
            for col in range(1, 19):
                ws.cell(row=row, column=col).fill = PAPER_FILL
        for col in range(2, 18):
            ws.cell(row=3, column=col).border = Border(bottom=INK_RULE)
        ws.row_dimensions[2].height = 32
        ws["B2"] = "Job search pipeline - dashboard"
        ws["B2"].font = TITLE_FONT
        ws["B2"].alignment = Alignment(vertical="center")
        ws["B3"] = (f"As of {self.d.as_of}. Tiles and the KPI table are LIVE formulas over "
                    "tblApplications; thresholds come from the Benchmarks sheet.")
        ws["B3"].font = SUBTITLE_FONT

        # --- KPI calculation table (right side): every tile reads a named cell here
        ws["K5"] = "KPI calculations (live)"
        ws["K5"].font = SECTION_FONT
        heads = ["Metric", "Value", "SQL metric (kpis.*)", "Defined name"]
        for j, h in enumerate(heads):
            c = ws.cell(row=6, column=11 + j, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.border = HEADER_BORDER
        ws.merge_cells("N6:O6")
        row = 7
        for kpi in KPIS:
            ws.cell(row=row, column=11, value=kpi.label).font = _font()
            cell = ws.cell(row=row, column=12, value=kpi.formula)
            cell.number_format = kpi.fmt
            cell.font = _font(10, True)
            cell.alignment = Alignment(horizontal="right")
            src = ("period_compare." if kpi.sql_column in ("last_30", "prev_30", "change") else
                   "kpis.") + kpi.sql_column if kpi.sql_column else "(Excel only)"
            ws.cell(row=row, column=13, value=src).font = _font(9, color=style.SECONDARY)
            ws.cell(row=row, column=14, value=f"kpi_{kpi.key}").font = _font(9, color=style.MUTED)
            for j in range(4):
                ws.cell(row=row, column=11 + j).border = Border(bottom=THIN)
            self._name(f"kpi_{kpi.key}", "Dashboard", f"$L${row}")
            row += 1
        kpi_end = row - 1

        # --- Tiles
        bands_txt = " / ".join(
            f"{b.label} {b.lower_pct:g}%+" for b in bm.RESPONSE_RATE_BANDS[:-1])
        tiles = [
            ("Applications", "=kpi_applications", FMT_INT,
             '="Last 30 days: "&kpi_last_30&"  (prev 30: "&kpi_prev_30&")"'),
            ("Response rate", "=kpi_response_rate", FMT_PCT,
             '=IF(kpi_applications>0,"95% CI "&TEXT(kpi_response_rate_lo,"0.0%")&" - "'
             '&TEXT(kpi_response_rate_hi,"0.0%"),"")'),
            ("Interview rate", "=kpi_interview_rate", FMT_PCT,
             '=IF(kpi_applications>0,TEXT(kpi_interview_vs_direct_pp,"+0.0;-0.0")'
             '&" pp vs "&TEXT(DirectInterviewBenchmark,"0.00")&"% direct benchmark","")'),
            ("Offers", "=kpi_offers", FMT_INT,
             '=IF(kpi_applications>0,"Offer rate "&TEXT(kpi_offer_rate,"0.0%"),"")'),
            ("Active pipeline", "=kpi_active", FMT_INT,
             '=kpi_needs_follow_up&" need a follow-up ("&FollowUpDays&"+ days quiet)"'),
            ("Median days to reply", "=kpi_median_days_to_response", FMT_DAYS,
             '="Known reply dates: "&kpi_known_reply_dates&" of "&kpi_responses&" replies"'),
            ("Response band", "=kpi_response_band", "General", f'="{bands_txt}"'),
            ("Direct channel share", "=kpi_direct_share", FMT_PCT,
             '=IF(kpi_applications>0,"Ghost rate "&TEXT(kpi_ghost_rate,"0.0%")&" | sample: "'
             '&kpi_sample_confidence,"")'),
        ]
        for idx, (label, formula, fmt, note) in enumerate(tiles):
            top = 5 if idx < 4 else 10
            c0 = 2 + (idx % 4) * 2
            self._tile(ws, top, c0, label, formula, fmt, note)

        # --- Funnel vs benchmark helper (live) + chart
        fr = kpi_end + 3
        self._section(ws, f"K{fr}", "Funnel per 100 applications vs published benchmark",
                      "Live: COUNTIFS over tblApplications; benchmark ranges from the Benchmarks sheet.")
        head = ["Stage", "You (per 100)", "Benchmark low", "Benchmark high", "Benchmark mid"]
        for j, h in enumerate(head):
            c = ws.cell(row=fr + 2, column=11 + j, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.border = HEADER_BORDER
        funnel = [("Employer responded", "kpi_responses", "BenchRespLo", "BenchRespHi"),
                  ("Assessment", "kpi_assessments", None, None),
                  ("Interview", "kpi_interviews", "BenchIntLo", "BenchIntHi"),
                  ("Offer", "kpi_offers", "BenchOfferLo", "BenchOfferHi")]
        for i, (stage, count, lo, hi) in enumerate(funnel):
            r = fr + 3 + i
            ws.cell(row=r, column=11, value=stage).font = _font()
            ws.cell(row=r, column=12,
                    value=f'=IF(kpi_applications>0,{count}/kpi_applications*100,0)').number_format = "0.0"
            if lo:
                ws.cell(row=r, column=13, value=f"={lo}").number_format = "0.0"
                ws.cell(row=r, column=14, value=f"={hi}").number_format = "0.0"
                ws.cell(row=r, column=15, value=f"=AVERAGE({lo},{hi})").number_format = "0.0"
        f_first, f_last = fr + 3, fr + 3 + len(funnel) - 1
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = "Funnel per 100 applications vs benchmark"
        chart.y_axis.title = "per 100 applications"
        cats = Reference(ws, min_col=11, min_row=f_first, max_row=f_last)
        for col in (12, 15):
            chart.add_data(Reference(ws, min_col=col, min_row=fr + 2, max_row=f_last),
                           titles_from_data=True)
        chart.set_categories(cats)
        self._colour_series(chart, [style.ACCENT, style.CONTEXT])
        self._chart_frame(chart, 10.9, 7.8)
        ws.add_chart(chart, "B15")

        # --- Platform response rates (live) + chart
        pr = f_last + 3
        self._section(ws, f"K{pr}", "Response rate by platform (live)",
                      "Indicative only at this sample size - read the CI and sample columns.")
        head = ["Platform", "Applications", "Responses", "Response rate", "CI low", "CI high",
                "Sample"]
        for j, h in enumerate(head):
            c = ws.cell(row=pr + 2, column=11 + j, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.border = HEADER_BORDER
        platforms = self.d.frames["platform_performance"]
        names = [str(p) for p in platforms["platform"].tolist()] if not platforms.empty else []
        self.platform_rows = names
        for i, platform in enumerate(names):
            r = pr + 3 + i
            ws.cell(row=r, column=11, value=platform).font = _font()
            ws.cell(row=r, column=12, value=f'=COUNTIFS({T}[platform],K{r})').number_format = FMT_INT
            ws.cell(row=r, column=13,
                    value=f'=COUNTIFS({T}[platform],K{r},{T}[is_response],TRUE)').number_format = FMT_INT
            ws.cell(row=r, column=14, value=f'=IF(L{r}>0,M{r}/L{r},"")').number_format = FMT_PCT
            for col, sign in ((15, "-"), (16, "+")):
                ws.cell(row=r, column=col, value=(
                    f"=IF(L{r}>0,MAX(0,MIN(1,(N{r}+WilsonZ^2/(2*L{r}){sign}WilsonZ*SQRT(N{r}*(1-N{r})/L{r}"
                    f"+WilsonZ^2/(4*L{r}^2)))/(1+WilsonZ^2/L{r}))),\"\")")).number_format = FMT_PCT
            ws.cell(row=r, column=17, value=(
                f'=IF(L{r}>=SampleAdequate,"Adequate",IF(L{r}>=SampleIndicative,"Indicative","Too small"))'))
        p_first, p_last = pr + 3, pr + 2 + max(1, len(names))
        for label, (bg, fg) in style.CONFIDENCE_CHIPS.items():
            ws.conditional_formatting.add(f"Q{p_first}:Q{p_last}", CellIsRule(
                operator="equal", formula=[f'"{label}"'], fill=_fill(bg), font=Font(color=fg)))
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Response rate by platform (live)"
        data = Reference(ws, min_col=14, min_row=pr + 2, max_row=p_last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=11, min_row=p_first, max_row=p_last))
        chart.x_axis.scaling.orientation = "maxMin"  # keep the table's order top-down
        chart.y_axis.number_format = "0%"
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        chart.legend = None
        self._colour_series(chart, [style.ACCENT])
        self._chart_frame(chart, 10.9, 7.8)
        ws.add_chart(chart, "F15")

        # --- Weekly velocity chart (SQL snapshot on the Weekly sheet)
        weekly = self.d.frames["weekly_velocity"]
        if not weekly.empty:
            n = len(weekly)
            wsw = self.sheets["Weekly"]
            bar = BarChart()
            bar.type = "col"
            bar.title = "Weekly applications with 4-week rolling average"
            bar.add_data(Reference(wsw, min_col=2, min_row=4, max_row=4 + n), titles_from_data=True)
            bar.set_categories(Reference(wsw, min_col=1, min_row=5, max_row=4 + n))
            self._colour_series(bar, [style.CONTEXT])
            line = LineChart()
            line.add_data(Reference(wsw, min_col=5, min_row=4, max_row=4 + n), titles_from_data=True)
            line.set_categories(Reference(wsw, min_col=1, min_row=5, max_row=4 + n))
            s = line.series[0]
            s.graphicalProperties.line.solidFill = style.ACCENT
            s.graphicalProperties.line.width = 28000
            s.smooth = False
            bar.x_axis.number_format = "dd mmm"
            bar.x_axis.tickLblSkip = 2
            bar.y_axis.title = "applications"
            bar += line
            self._chart_frame(bar, 21.8, 7.8)
            ws.add_chart(bar, "B31")
        ws.freeze_panes = None
        self._print(ws)

    def _tile(self, ws, top: int, c0: int, label: str, formula: str, fmt: str, note: str) -> None:
        for r in range(top, top + 4):
            for c in (c0, c0 + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = SURFACE_FILL
                cell.border = Border(
                    left=THIN if c == c0 else None, right=THIN if c == c0 + 1 else None,
                    top=ACCENT_RULE if r == top else None, bottom=THIN if r == top + 3 else None)
        a, b = get_column_letter(c0), get_column_letter(c0 + 1)
        ws.merge_cells(f"{a}{top}:{b}{top}")
        ws.merge_cells(f"{a}{top + 1}:{b}{top + 2}")
        ws.merge_cells(f"{a}{top + 3}:{b}{top + 3}")
        lab = ws[f"{a}{top}"]
        lab.value = label.upper()
        lab.font = _font(8, True, style.SECONDARY)
        lab.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        val = ws[f"{a}{top + 1}"]
        val.value = formula
        val.number_format = fmt
        val.font = _font(22, True, style.ACCENT if top == 5 else style.INK, name=HEAD)
        val.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        sub = ws[f"{a}{top + 3}"]
        sub.value = note
        sub.font = _font(8, False, style.MUTED)
        sub.alignment = Alignment(horizontal="left", vertical="top", indent=1, shrink_to_fit=True)
        for r in range(top, top + 4):
            ws.row_dimensions[r].height = 17 if r != top + 3 else 20

    @staticmethod
    def _colour_series(chart, colours: list[str]) -> None:
        for series, colour in zip(chart.series, colours, strict=False):
            series.graphicalProperties.solidFill = colour
            series.graphicalProperties.line.solidFill = colour

    @staticmethod
    def _chart_frame(chart, width_cm: float, height_cm: float) -> None:
        """Editorial chart chrome: Georgia title, grey axis labels, hairline gridlines."""
        chart.width = width_cm
        chart.height = height_cm
        chart.style = 2
        title = _title_text(chart.title)
        if title:
            cp = CharacterProperties(sz=1200, b=True, solidFill=style.INK,
                                     latin=DrawingFont(typeface=HEAD))
            para = Paragraph(pPr=ParagraphProperties(defRPr=cp),
                             r=[RegularTextRun(rPr=cp, t=title)])
            chart.title = Title(tx=Text(rich=RichText(p=[para])), overlay=False)
        for axis in (chart.x_axis, chart.y_axis):
            axis.delete = False
            axis.txPr = _rich(8, style.SECONDARY)
            axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=style.RULE))
        chart.y_axis.majorGridlines = ChartLines(
            spPr=GraphicalProperties(ln=LineProperties(solidFill=style.RULE, w=6350)))
        chart.x_axis.majorGridlines = None
        if chart.legend is not None:
            chart.legend.position = "b"
            chart.legend.txPr = _rich(8, style.SECONDARY)
        chart.graphical_properties = GraphicalProperties(
            solidFill=style.SURFACE, ln=LineProperties(solidFill=style.RULE, w=6350))

    def _report(self, sheet: str, title: str, query: str) -> None:
        ws = self.sheets[sheet]
        comment = " ".join(line.strip().lstrip("-").strip() for line in SQL[query].splitlines()
                           if line.strip().startswith("--"))
        self._banner(ws, title, f"SQL snapshot of metrics.{query} as of {self.d.as_of}. {comment}", 12)
        ws["A2"].alignment = Alignment(wrap_text=False)

    # ---- Funnel -----------------------------------------------------------------------
    def _funnel(self) -> None:
        ws = self.sheets["Funnel"]
        self._report("Funnel", "Application funnel vs benchmark", "funnel")
        frame = self.d.frames["funnel"]
        first, last, _, _ = self._write_table(
            ws, frame, 4, 1, "tblFunnel",
            formats={"pct_of_applied": FMT_PCT_NUM, "step_conversion_pct": FMT_PCT_NUM,
                     "benchmark_per_100_lo": "0.0", "benchmark_per_100_hi": "0.0"},
            widths={"stage": 22, "vs_benchmark": 14})
        for label, colour in (("Ahead", style.POSITIVE), ("On par", style.CAUTION),
                              ("Behind", style.NEGATIVE)):
            ws.conditional_formatting.add(f"J{first}:J{last}", CellIsRule(
                operator="equal", formula=[f'"{label}"'], font=Font(color=colour, bold=True)))
        ws.conditional_formatting.add(f"D{first}:D{last}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=100,
            color=style.STATUS["Assessment / OA"]))
        ws.cell(row=last + 2, column=1, value=(
            "Furthest stage reached: an application rejected after its interview still counts as an "
            "interview. pct_of_applied = per 100 applications, the unit of the published benchmark. "
            "Percent columns are in percent units (25.0 = 25%).")).font = NOTE_FONT
        ws.freeze_panes = "A5"
        self._print(ws)

    # ---- Platforms ---------------------------------------------------------------------
    def _platforms(self) -> None:
        ws = self.sheets["Platforms"]
        self._report("Platforms", "Platform performance", "platform_performance")
        frame = self.d.frames["platform_performance"].copy()
        pct_cols = [c for c in frame.columns if c.endswith("_pct") or c.endswith("_lo")
                    or c.endswith("_hi") or c.endswith("_pp")]
        first, last, _, last_col = self._write_table(
            ws, frame, 4, 1, "tblPlatforms",
            formats={c: FMT_PCT_NUM for c in pct_cols} | {"median_days_to_response": FMT_DAYS},
            widths={"platform": 16, "channel": 22, "sample_confidence": 18})
        cols = list(frame.columns)
        col = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
        if "response_rate_pct" in col:
            ws.conditional_formatting.add(
                f"{col['response_rate_pct']}{first}:{col['response_rate_pct']}{last}",
                DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100,
                            color=style.STATUS["Assessment / OA"]))
        if "sample_confidence" in col:
            rng = f"{col['sample_confidence']}{first}:{col['sample_confidence']}{last}"
            for label, (bg, fg) in style.CONFIDENCE_CHIPS.items():
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=[f'"{label}"'], fill=_fill(bg), font=Font(color=fg)))
        if "vs_direct_benchmark_pp" in col:
            rng = f"{col['vs_direct_benchmark_pp']}{first}:{col['vs_direct_benchmark_pp']}{last}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                          font=Font(color=style.NEGATIVE)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual",
                                                          formula=["0"], font=Font(color=style.POSITIVE)))
        note_row = last + 2
        ws.cell(row=note_row, column=1, value=(
            f"Wilson score 95% intervals (z = {Z}). Sample confidence: Adequate n >= "
            f"{bm.SAMPLE_ADEQUATE}, Indicative n >= {bm.SAMPLE_INDICATIVE}, else Too small. "
            f"vs_direct_benchmark_pp = interview rate minus the {bm.DIRECT_INTERVIEW_RATE_PCT}% "
            "direct-application benchmark. Platform differences are indicative, not statistically "
            "significant at this sample size.")).font = NOTE_FONT

        # Chart helper: CI error-bar lengths (formulas over the table) + chart
        if not frame.empty and {"response_rate_pct", "response_rate_lo", "response_rate_hi"} <= set(cols):
            hc = last_col + 2
            h1, h2 = get_column_letter(hc), get_column_letter(hc + 1)
            ws.cell(row=4, column=hc, value="err_minus (chart)").font = NOTE_FONT
            ws.cell(row=4, column=hc + 1, value="err_plus (chart)").font = NOTE_FONT
            for r in range(first, last + 1):
                ws[f"{h1}{r}"] = f"={col['response_rate_pct']}{r}-{col['response_rate_lo']}{r}"
                ws[f"{h2}{r}"] = f"={col['response_rate_hi']}{r}-{col['response_rate_pct']}{r}"
                ws[f"{h1}{r}"].number_format = ws[f"{h2}{r}"].number_format = FMT_PCT_NUM
            ws.column_dimensions[h1].width = ws.column_dimensions[h2].width = 15
            chart = BarChart()
            chart.type = "col"
            chart.title = "Response rate by platform with 95% Wilson interval"
            rr = cols.index("response_rate_pct") + 1
            chart.add_data(Reference(ws, min_col=rr, min_row=4, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
            series = chart.series[0]
            series.errBars = ErrorBars(
                errDir="y", errBarType="both", errValType="cust", noEndCap=False,
                minus=NumDataSource(numRef=NumRef(f=f"'Platforms'!${h1}${first}:${h1}${last}")),
                plus=NumDataSource(numRef=NumRef(f=f"'Platforms'!${h2}${first}:${h2}${last}")))
            chart.y_axis.title = "response rate (%)"
            chart.legend = None
            self._colour_series(chart, [style.ACCENT])
            self._chart_frame(chart, 22, 8.5)
            ws.add_chart(chart, f"A{note_row + 2}")
        ws.freeze_panes = "B5"
        self._print(ws)

    # ---- Weekly ---------------------------------------------------------------------------
    def _weekly(self) -> None:
        ws = self.sheets["Weekly"]
        self._report("Weekly", "Weekly velocity", "weekly_velocity")
        frame = self.d.frames["weekly_velocity"]
        first, last, _, _ = self._write_table(
            ws, frame, 4, 1, "tblWeekly",
            formats={"week_start": FMT_DATE, "rolling_4w_avg": "0.0"},
            widths={"week_start": 13})
        if not frame.empty:
            ws.conditional_formatting.add(f"B{first}:B{last}", DataBarRule(
                start_type="num", start_value=0, end_type="max", color=style.CONTEXT_DARK))
            chart = LineChart()
            chart.title = "Cumulative applications"
            chart.add_data(Reference(ws, min_col=6, min_row=4, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
            chart.x_axis.number_format = "dd mmm"
            chart.legend = None
            self._colour_series(chart, [style.ACCENT])
            self._chart_frame(chart, 18, 8)
            ws.add_chart(chart, "H4")
        ws.freeze_panes = "A5"
        self._print(ws)

    # ---- Cohorts ----------------------------------------------------------------------------
    def _cohorts(self) -> None:
        ws = self.sheets["Cohorts"]
        self._report("Cohorts", "Monthly application cohorts", "monthly_cohorts")
        frame = self.d.frames["monthly_cohorts"]
        first, last, _, _ = self._write_table(
            ws, frame, 4, 1, "tblCohorts",
            formats={"response_rate_pct": FMT_PCT_NUM, "interview_rate_pct": FMT_PCT_NUM},
            widths={"cohort_month": 14})
        cols = list(frame.columns)
        for name in ("response_rate_pct", "interview_rate_pct"):
            if name in cols:
                letter = get_column_letter(cols.index(name) + 1)
                ws.conditional_formatting.add(f"{letter}{first}:{letter}{last}", ColorScaleRule(
                    start_type="min", start_color=style.SURFACE, end_type="max",
                    end_color=style.tint(style.ACCENT, 0.55)))
        if "is_maturing" in cols:
            letter = get_column_letter(cols.index("is_maturing") + 1)
            ws.conditional_formatting.add(f"A{first}:{get_column_letter(len(cols))}{last}",
                                          FormulaRule(formula=[f"${letter}{first}=TRUE"],
                                                      font=Font(italic=True, color=style.MUTED)))
        ws.cell(row=last + 2, column=1, value=(
            "Cohort = month the application was sent. Cohorts younger than 30 days (is_maturing, "
            "shown in grey italics) are still collecting replies - read their rates as floors.")
                ).font = NOTE_FONT
        ws.freeze_panes = "B5"
        self._print(ws)

    # ---- Reply time -------------------------------------------------------------------------
    def _reply_time(self) -> None:
        ws = self.sheets["Reply time"]
        self._report("Reply time", "Time to first reply (Kaplan-Meier)", "response_survival")
        frame = self.d.frames["response_survival"]
        first, last, _, _ = self._write_table(
            ws, frame, 4, 1, "tblSurvival",
            formats={"still_waiting_pct": "0.00", "replied_pct": "0.00"})
        if not frame.empty:
            chart = ScatterChart()
            chart.title = "Share of applications with a reply, by days since applying"
            chart.style = 2
            xs = Reference(ws, min_col=1, min_row=first, max_row=last)
            ys = Reference(ws, min_col=6, min_row=4, max_row=last)
            series = Series(ys, xs, title_from_data=True)
            series.graphicalProperties.line.solidFill = style.ACCENT
            series.graphicalProperties.line.width = 28000
            series.marker.symbol = "none"
            series.smooth = False
            chart.series.append(series)
            chart.x_axis.title = "days since applying"
            chart.y_axis.title = "replied (%)"
            chart.legend = None
            self._chart_frame(chart, 17, 8)
            ws.add_chart(chart, "H4")
        r = max(last + 3, 22)
        self._section(ws, f"H{r}", "Reply-time buckets (known reply dates only)")
        buckets = self.d.frames["response_time_buckets"]
        _, b_last, _, _ = self._write_table(ws, buckets, r + 1, 8, "tblReplyBuckets",
                                            formats={"share_pct": FMT_PCT_NUM},
                                            widths={"bucket": 12})
        r2 = b_last + 3
        self._section(ws, f"H{r2}", "Typical time to each milestone (medians, known dates only)")
        self._write_table(ws, self.d.frames["stage_timing"], r2 + 1, 8, "tblStageTiming",
                          formats={"median_days": FMT_DAYS, "p25_days": FMT_DAYS,
                                   "p75_days": FMT_DAYS}, widths={"milestone": 14})
        ws.cell(row=last + 2, column=1, value=(
            "Kaplan-Meier: applications still waiting are right-censored at their current age, not "
            "dropped - dropping them would make employers look faster than they are. Replies with "
            "an unknown date are excluded (they cannot be placed on the time axis).")).font = NOTE_FONT
        ws.freeze_panes = "A5"
        self._print(ws)

    # ---- Follow-ups ---------------------------------------------------------------------
    def _follow_ups(self) -> None:
        ws = self.sheets["Follow-ups"]
        self._report("Follow-ups", "Follow-up queue", "follow_up_queue")
        ws["A2"].value = (f"Open applications silent for {self.d.follow_up_days}+ days as of "
                          f"{self.d.as_of} (SQL snapshot of metrics.follow_up_queue), oldest "
                          "silence first.")
        frame = self.d.frames["follow_up_queue"]
        first, last, _, _ = self._write_table(
            ws, frame, 4, 1, "tblFollowUps",
            formats={"applied_date": FMT_DATE, "last_activity_date": FMT_DATE},
            widths={"application_id": 20, "company_name": 26, "role_title": 32, "gmail_link": 28,
                    "job_url": 36, "suggested_action": 34, "current_status": 16})
        cols = list(frame.columns)
        for name in ("gmail_link", "job_url"):
            if name in cols:
                letter = get_column_letter(cols.index(name) + 1)
                for r in range(first, last + 1):
                    cell = ws[f"{letter}{r}"]
                    if isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value
                        cell.font = LINK_FONT
        if "days_since_activity" in cols and not frame.empty:
            letter = get_column_letter(cols.index("days_since_activity") + 1)
            ws.conditional_formatting.add(f"{letter}{first}:{letter}{last}", ColorScaleRule(
                start_type="num", start_value=self.d.follow_up_days,
                start_color=style.tint(style.CAUTION, 0.75),
                end_type="num", end_value=self.d.ghost_after_days,
                end_color=style.tint(style.NEGATIVE, 0.65)))
        if frame.empty:
            ws.cell(row=first + 1, column=1, value="Nothing to chase - every open application "
                    "has had activity recently.").font = NOTE_FONT
        ws.freeze_panes = "C5"
        self._print(ws, title_rows="4:4")

    # ---- Segments ----------------------------------------------------------------------
    def _segments(self) -> None:
        ws = self.sheets["Segments"]
        self._banner(ws, "Segment performance",
                     f"SQL snapshot of metrics.segment_performance as of {self.d.as_of}: role "
                     "family, industry, seniority and job type, with Wilson 95% intervals.", 10)
        row = 4
        for seg, title in (("role_family", "By role family"), ("industry", "By industry"),
                           ("seniority", "By seniority"), ("job_type", "By job type")):
            self._section(ws, f"A{row}", title)
            frame = self.d.frames[f"segment_{seg}"]
            pct = {c: FMT_PCT_NUM for c in frame.columns if c.endswith(("_pct", "_lo", "_hi"))}
            name = "tblSeg" + "".join(p.capitalize() for p in seg.split("_"))
            _, last, _, _ = self._write_table(ws, frame, row + 1, 1, name,
                                              formats=pct | {"median_days_to_response": FMT_DAYS},
                                              widths={"segment": 30})
            row = last + 3
        ws.cell(row=row - 1, column=1, value=(
            "Segments with few applications are indicative only - compare the CI columns, not the "
            "point estimates.")).font = NOTE_FONT
        self._print(ws)

    # ---- Data quality -------------------------------------------------------------------
    def _data_quality(self) -> None:
        ws = self.sheets["Data quality"]
        self._report("Data quality", "Data quality audit", "data_quality")
        frame = self.d.frames["data_quality"].copy()
        frame["coverage"] = None
        first, last, _, _ = self._write_table(ws, frame, 4, 1, "tblDataQuality",
                                              widths={"check_name": 34, "coverage": 12})
        for r in range(first, last + 1):
            ws[f"D{r}"] = f'=IF(C{r}>0,B{r}/C{r},"")'
            ws[f"D{r}"].number_format = FMT_PCT
        ws.conditional_formatting.add(f"D{first}:D{last}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color=style.POSITIVE))
        ws.cell(row=last + 1, column=1, value=(
            "coverage is a live formula (passing / total). 'Reply date known' is measured against "
            "replies: replies without a trustworthy date keep days_to_response blank.")
                ).font = NOTE_FONT
        r = last + 4
        self._section(ws, f"A{r}", "Where every application stands now (metrics.status_mix)")
        mix = self.d.frames["status_mix"]
        m_first, m_last, _, _ = self._write_table(
            ws, mix, r + 1, 1, "tblStatusMix",
            formats={"share_pct": FMT_PCT_NUM, "avg_days_since_activity": FMT_DAYS})
        for status, (bg, fg) in style.STATUS_CHIPS.items():
            ws.conditional_formatting.add(f"A{m_first}:A{m_last}", CellIsRule(
                operator="equal", formula=[f'"{status}"'], fill=_fill(bg), font=Font(color=fg)))
        ws.freeze_panes = "A5"
        self._print(ws)

    # ---- Data dictionary ---------------------------------------------------------------
    def _dictionary(self) -> None:
        ws = self.sheets["Data dictionary"]
        self._banner(ws, "Data dictionary",
                     "One metric, four tools: each metric is defined once in SQL and reproduced "
                     "exactly in Excel, DAX and Tableau. Then every column of tblApplications.", 6)
        dax = {m.name: m.expression for m in dax_measure_list(self.d.follow_up_days)}
        tab = {c.name: c.formula for c in tableau_calc_list()}
        self._section(ws, "A4", "Metric definitions - one metric, four tools",
                      "Excel = the live formula in the Dashboard KPI table (defined names kpi_*, "
                      "FollowUpDays, WilsonZ ... point at Dashboard / Benchmarks cells).")
        rows = []
        for kpi in KPIS:
            if not kpi.sql:
                continue
            rows.append([kpi.label, kpi.definition, kpi.sql, kpi.formula,
                         (f"[{kpi.dax}] =\n{dax[kpi.dax]}" if kpi.dax in dax else ""),
                         (f"[{kpi.tableau}] =\n{tab[kpi.tableau]}" if kpi.tableau in tab else "")])
        frame = pd.DataFrame(rows, columns=["Metric", "Definition", "SQL (DuckDB)", "Excel",
                                            "DAX (Power BI)", "Tableau"])
        first, last, _, _ = self._write_table(
            ws, frame, 6, 1, "tblMetricDefinitions",
            widths={"Metric": 26, "Definition": 44, "SQL (DuckDB)": 58, "Excel": 58,
                    "DAX (Power BI)": 58, "Tableau": 58})
        mono = Font(name="Consolas", size=9, color=style.INK)
        for r in range(first, last + 1):
            for c in range(1, 7):
                cell = ws.cell(row=r, column=c)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if c >= 3:
                    cell.font = mono
                if c == 4:
                    _as_text(cell, str(cell.value))
            ws.cell(row=r, column=1).font = _font(10, True)
            lines = max(str(ws.cell(row=r, column=c).value or "").count("\n") + 1 +
                        len(str(ws.cell(row=r, column=c).value or "")) // 60 for c in range(2, 7))
            ws.row_dimensions[r].height = min(240, 13 * max(2, lines))

        top = last + 3
        self._section(ws, f"A{top}", "tblApplications columns",
                      "Same names as job_applications.csv (Tableau) and fct_applications (Power BI).")
        dict_rows = []
        for col in self.app_cols:
            kind = catalog.column_type(col)
            fmt = {"date": "date (yyyy-mm-dd)", "boolean": "TRUE / FALSE", "integer": "whole number",
                   "decimal": "decimal (0.00)"}.get(kind, "text")
            dict_rows.append([col, catalog.description(col), fmt])
        dframe = pd.DataFrame(dict_rows, columns=["Column", "Description", "Type"])
        d_first, d_last, _, _ = self._write_table(ws, dframe, top + 2, 1, "tblDictionary")
        for r in range(d_first, d_last + 1):
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=1).font = Font(name="Consolas", size=9, bold=True)
        ws.cell(row=d_last + 2, column=1, value=(
            "Surrogate keys (*_key) are in the star-schema CSVs for Power BI; they are left out of "
            "tblApplications.")).font = NOTE_FONT
        ws.freeze_panes = "B7"
        self._print(ws)

    # ---- Cover ---------------------------------------------------------------------------
    def _cover(self) -> None:
        ws = self.sheets["Cover"]
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 70
        ws.column_dimensions["D"].width = 26
        for col in range(2, 5):
            ws.cell(row=2, column=col).border = Border(top=Side(style="thick", color=style.ACCENT))
            ws.cell(row=5, column=col).border = Border(bottom=INK_RULE)
        ws.row_dimensions[3].height = 38
        ws["B3"] = "Job Search Analytics"
        ws["B3"].font = _font(28, True, style.INK, name=HEAD)
        ws["B4"] = ("Application pipeline: Gmail -> deterministic parser -> dbt / DuckDB star "
                    "schema -> SQL metric layer -> Excel, Power BI, Tableau")
        ws["B4"].font = SUBTITLE_FONT
        k = self.d.kpis
        first = _py(k.get("first_applied"))
        last = _py(k.get("last_applied"))
        facts = [
            ("As of", f"{self.d.built_at} (warehouse build)"),
            ("Applications", f"{int(k.get('applications') or 0)} to "
                             f"{int(k.get('companies') or 0)} companies"),
            ("Period", f"{first} to {last}" if first else "no applications yet"),
        ]
        r = 7
        for label, value in facts:
            ws.cell(row=r, column=2, value=label).font = _font(10, True, style.INK)
            ws.cell(row=r, column=3, value=value).font = _font()
            r += 1

        r += 1
        self._section(ws, f"B{r}", "Contents")
        r += 1
        heads = ["Sheet", "What it shows", "Numbers are"]
        for j, h in enumerate(heads):
            c = ws.cell(row=r, column=2 + j, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.border = HEADER_BORDER
        contents = [
            ("Dashboard", "KPI tiles, funnel vs benchmark, platform response rates, weekly velocity",
             "Live formulas"),
            ("Applications", "tblApplications - one row per application, filterable", "Data"),
            ("Funnel", "Stage counts, LAG step conversion, published per-100 benchmark", "SQL snapshot"),
            ("Platforms", "Response / interview rate per platform, Wilson 95% CI, sample confidence",
             "SQL snapshot"),
            ("Weekly", "Applications per ISO week, 4-week rolling average, running total",
             "SQL snapshot"),
            ("Cohorts", "Month-of-application cohorts and how each has played out", "SQL snapshot"),
            ("Reply time", "Kaplan-Meier time-to-first-reply, reply buckets, stage timing",
             "SQL snapshot"),
            ("Follow-ups", "Open applications gone quiet, with a suggested action", "SQL snapshot"),
            ("Segments", "Role family, industry, seniority and job type performance", "SQL snapshot"),
            ("Data quality", "Coverage audit of every attribute the analysis relies on, status mix",
             "SQL snapshot + formulas"),
            ("Data dictionary", "Metric definitions in SQL / Excel / DAX / Tableau; every column",
             "Reference"),
            ("Benchmarks", "Published hiring benchmarks, response bands, follow-up setting",
             "Inputs"),
        ]
        for sheet, what, kind in contents:
            r += 1
            cell = ws.cell(row=r, column=2, value=sheet)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1", display=sheet)
            cell.font = LINK_FONT
            ws.cell(row=r, column=3, value=what).font = _font()
            ws.cell(row=r, column=4, value=kind).font = _font(9, color=style.SECONDARY)
            for c in range(2, 5):
                ws.cell(row=r, column=c).border = Border(bottom=THIN)

        r += 2
        self._section(ws, f"B{r}", "Read before quoting a number")
        known = int(k.get("responses_with_known_date") or 0)
        responses = int(k.get("responses") or 0)
        notes = [
            ("Reply timing", "days_to_response is BLANK when the reply date is unknown and is never "
             f"zero-filled. Medians use known dates only ({known} of {responses} replies)."),
            ("Small samples", "Rates carry Wilson score 95% intervals. Platform and segment "
             "comparisons are indicative, not statistically significant at this sample size - "
             f"check the sample label (Adequate n >= {bm.SAMPLE_ADEQUATE}, Indicative n >= "
             f"{bm.SAMPLE_INDICATIVE})."),
            ("Benchmarks", "Published aggregate figures (mixed seniority, mostly cold applications): "
             "context, not targets. Response bands: " + ", ".join(
                 f"{b.label} {b.lower_pct:g}%+" for b in bm.RESPONSE_RATE_BANDS[:-1])
             + f"; direct-application interview benchmark {bm.DIRECT_INTERVIEW_RATE_PCT}%."),
            ("Funnel", "Stages use the furthest stage reached: a rejection after an interview "
             "still counts as an interview."),
            ("Snapshot", "Day counts (days_since_applied, days_since_activity) are as of the "
             "warehouse build; re-export to roll them forward."),
            ("Scope", "This is job-application pipeline data (one row per application). Insurance "
             "portfolio measures do not apply to it and are deliberately absent."),
        ]
        for label, text in notes:
            r += 1
            ws.cell(row=r, column=2, value=label).font = _font(10, True, style.INK)
            cell = ws.cell(row=r, column=3, value=text)
            cell.font = _font()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
            ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(text) / 75))

        r += 2
        self._section(ws, f"B{r}", "Refreshing")
        for label, text in (
            ("Full refresh", "python -m jobtracker export - re-runs the SQL and rebuilds every sheet."),
            ("Quick refresh", "Paste new rows from job_applications.csv into tblApplications. The "
             "Dashboard formulas recalculate; the SQL snapshot sheets keep their export date."),
            ("Follow-up window", "Edit FollowUpDays on the Benchmarks sheet (the blue input cell)."),
        ):
            r += 1
            ws.cell(row=r, column=2, value=label).font = _font(10, True, style.INK)
            ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
            cell = ws.cell(row=r, column=3, value=text)
            cell.font = _font()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(text) / 75))
        self._print(ws, landscape=False)


def build_workbook(settings) -> bytes:
    """Build the analyst workbook from the warehouse and return the .xlsx bytes."""
    from ..warehouse import connect

    with connect(settings) as con:
        data = _collect(con, settings)
    wb = _Builder(data).build()
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def kpi_table_rows() -> list[tuple[str, str, float]]:
    """(label, SQL column, scale) for every Dashboard KPI that has a SQL twin - used by the tests."""
    return [(k.label, k.sql_column, k.scale) for k in KPIS if k.sql_column]


__all__ = ["KPIS", "SHEETS", "build_workbook", "kpi_table_rows"]
