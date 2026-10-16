"""Prove the workbook's live formulas equal the SQL metric layer.

``recalculate`` opens the workbook in headless LibreOffice, which computes every formula on
load, and saves a copy with cached values. ``compare_dashboard`` reads the Dashboard KPI table
back and compares each value with ``Metrics.kpis()`` / ``period_compare`` (rates rounded to one
decimal, exactly like the SQL). Used by ``tests/test_bi.py``; also handy by hand:

    python -m jobtracker.bi.verify --data-dir <dir>
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
import warnings
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore", message=".*extension is not supported.*", module="openpyxl")


def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalculate(xlsx: Path, out_dir: Path, timeout: int = 120) -> Path:
    """Recalculate ``xlsx`` with LibreOffice; returns the path of the recalculated copy."""
    binary = soffice_path()
    if not binary:
        raise FileNotFoundError("LibreOffice (soffice) is not installed")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
        env = {**os.environ, "SAL_USE_VCLPLUGIN": "svp"}
        subprocess.run(
            [binary, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless",
             "--norestore", "--convert-to", "xlsx", "--outdir", str(out_dir), str(xlsx)],
            check=True, capture_output=True, timeout=timeout, env=env)
    result = out_dir / xlsx.name
    if not result.exists():
        raise RuntimeError(f"LibreOffice did not write {result}")
    return result


def read_kpi_table(path: Path) -> dict[str, object]:
    """{SQL metric label (e.g. 'kpis.response_rate_pct'): cached value} from the Dashboard."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Dashboard"]
    out: dict[str, object] = {}
    for row in ws.iter_rows(min_row=7, min_col=11, max_col=13):
        label, value, source = (c.value for c in row)
        if not label or not source or not str(source).startswith(("kpis.", "period_compare.")):
            if label is None:
                break
            continue
        out[str(source)] = value
    wb.close()
    return out


def read_platform_block(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Dashboard"]
    out: dict[str, dict] = {}
    start = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=11).value == "Platform" and ws.cell(row=r, column=12).value == "Applications":
            start = r + 1
            break
    if start:
        r = start
        while ws.cell(row=r, column=11).value:
            vals = [ws.cell(row=r, column=c).value for c in range(11, 18)]
            out[str(vals[0])] = dict(zip(("applications", "responses", "rate", "lo", "hi",
                                          "sample"), vals[1:], strict=True))
            r += 1
    wb.close()
    return out


def _norm(value):
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "to_pydatetime"):
        return value.to_pydatetime().date()
    if hasattr(value, "item"):
        return value.item()
    return value


def compare_dashboard(path: Path, settings) -> list[str]:
    """Return a list of mismatches between the recalculated Dashboard and the SQL (empty = OK)."""
    from ..warehouse import connect
    from ..warehouse.metrics import Metrics
    from .workbook import kpi_table_rows

    with connect(settings) as con:
        metrics = Metrics(con, settings)
        sql = {f"kpis.{k}": v for k, v in metrics.kpis().items()}
        period = metrics.query("period_compare").iloc[0].to_dict()
        sql.update({f"period_compare.{k}": v for k, v in period.items()})
        platforms = metrics.query("platform_performance")
    excel = read_kpi_table(path)
    problems: list[str] = []
    for label, column, scale in kpi_table_rows():
        key = ("period_compare." if column in ("last_30", "prev_30", "change") else "kpis.") + column
        want = _norm(sql.get(key))
        got = _norm(excel.get(key))
        if want is None or (isinstance(want, float) and want != want):
            if got not in (None, ""):
                problems.append(f"{label}: SQL is NULL, Excel {got!r}")
            continue
        if isinstance(want, (date, str)):
            if got != want:
                problems.append(f"{label}: SQL {want!r} != Excel {got!r}")
            continue
        if not isinstance(got, (int, float)):
            problems.append(f"{label}: Excel returned {got!r} (SQL {want!r})")
            continue
        value = float(got) * scale
        if scale != 1:
            value = round(value + 1e-9, 1)
        if abs(value - float(want)) > 0.051:
            problems.append(f"{label}: SQL {want} != Excel {value}")
    block = read_platform_block(path)
    for row in platforms.itertuples():
        got = block.get(str(row.platform))
        if not got:
            problems.append(f"platform {row.platform}: missing from the Dashboard")
            continue
        checks = [("applications", row.applications, 1), ("responses", row.responses, 1),
                  ("rate", row.response_rate_pct, 100), ("lo", row.response_rate_lo, 100),
                  ("hi", row.response_rate_hi, 100)]
        for name, want, scale in checks:
            value = float(got[name]) * scale
            if abs(round(value + 1e-9, 1) - float(want)) > 0.051:
                problems.append(f"platform {row.platform} {name}: SQL {want} != Excel {value}")
        if got["sample"] != row.sample_confidence:
            problems.append(f"platform {row.platform} sample: {got['sample']} != {row.sample_confidence}")
    return problems


def main(argv: list[str] | None = None) -> int:
    import argparse

    from ..config import load_settings
    from .workbook import build_workbook

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir")
    args = parser.parse_args(argv)
    settings = load_settings(args.data_dir)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "workbook.xlsx"
        src.write_bytes(build_workbook(settings))
        recalculated = recalculate(src, Path(tmp) / "recalc")
        values = read_kpi_table(recalculated)
        problems = compare_dashboard(recalculated, settings)
    for key, value in values.items():
        print(f"{key:<40} {value}")
    print("OK - Excel formulas match SQL" if not problems else "\n".join(problems))
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
