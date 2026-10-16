"""BI pack: workbook structure, Excel-formula / SQL parity (LibreOffice), DAX + Tableau field
references against the live warehouse, benchmark thresholds and honesty rules.

Hermetic: builds the synthetic demo dataset into a temp dir (no network, no personal data).
"""
from __future__ import annotations

import io
import json
import re
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from jobtracker import benchmarks as bm
from jobtracker.bi import (
    build_workbook,
    catalog,
    dax_measures,
    export_all,
    model_spec,
    powerbi_theme,
    tableau_calculations,
    verify,
)
from jobtracker.bi.dax import render_dax_file
from jobtracker.bi.tableau import PARAMETERS
from jobtracker.config import load_settings
from jobtracker.demo import build_demo
from jobtracker.warehouse import connect
from jobtracker.warehouse.metrics import Metrics, Z

ROOT = Path(__file__).resolve().parent.parent
INSURANCE_TERMS = re.compile(r"loss[\s_-]*ratio|\bibnr\b|\bcsm\b|premium|\bclaims?\b", re.I)
_tmp: tempfile.TemporaryDirectory | None = None
SETTINGS = None
WORKBOOK: bytes = b""


def setUpModule():
    global _tmp, SETTINGS, WORKBOOK
    _tmp = tempfile.TemporaryDirectory(prefix="bi-test-")
    data_dir = Path(_tmp.name) / "demo"
    build_demo(data_dir, seed=7)
    SETTINGS = load_settings(data_dir)
    WORKBOOK = build_workbook(SETTINGS)


def tearDownModule():
    if _tmp is not None:
        _tmp.cleanup()


def _warehouse_columns() -> dict[str, set[str]]:
    with connect(SETTINGS) as con:
        rows = con.execute(
            "SELECT table_name, column_name FROM information_schema.columns").fetchall()
    out: dict[str, set[str]] = {}
    for table, column in rows:
        out.setdefault(table, set()).add(column)
    return out


def _wb(data_only: bool = False):
    return load_workbook(io.BytesIO(WORKBOOK), data_only=data_only)


class WorkbookStructureTest(unittest.TestCase):
    def test_sheets_and_tables(self):
        wb = _wb()
        for sheet in ("Cover", "Dashboard", "Applications", "Funnel", "Platforms", "Weekly",
                      "Cohorts", "Reply time", "Follow-ups", "Data quality", "Data dictionary",
                      "Benchmarks"):
            self.assertIn(sheet, wb.sheetnames)
        self.assertEqual(wb.sheetnames[0], "Cover")
        tables = {name for ws in wb for name in ws.tables}
        for name in ("tblApplications", "Benchmarks", "ResponseBands", "tblPlatforms",
                     "tblFunnel", "tblWeekly", "tblCohorts", "tblSurvival", "tblFollowUps",
                     "tblDataQuality", "tblMetricDefinitions", "tblDictionary"):
            self.assertIn(name, tables)
        self.assertFalse(wb["Dashboard"].sheet_view.showGridLines)

    def test_applications_table_matches_fact_rows(self):
        wb = _wb()
        ws = wb["Applications"]
        ref = ws.tables["tblApplications"].ref
        first, last = ref.split(":")
        n_rows = int(re.sub(r"\D", "", last)) - int(re.sub(r"\D", "", first))
        with connect(SETTINGS) as con:
            facts = con.execute("SELECT COUNT(*) FROM fct_applications").fetchone()[0]
        self.assertGreater(facts, 0)
        self.assertEqual(n_rows, facts)
        headers = [c.value for c in ws[1]]
        self.assertFalse([h for h in headers if str(h).endswith("_key")])
        self.assertEqual(ws.freeze_panes, "D2")
        self.assertTrue(ws.conditional_formatting)

    def test_unknown_reply_times_stay_blank(self):
        wb = _wb()
        ws = wb["Applications"]
        headers = [c.value for c in ws[1]]
        col = headers.index("days_to_response")
        values = [row[col] for row in ws.iter_rows(min_row=2, values_only=True)]
        with connect(SETTINGS) as con:
            known, nulls = con.execute(
                "SELECT COUNT(days_to_response), COUNT(*) - COUNT(days_to_response) "
                "FROM fct_applications").fetchone()
        self.assertEqual(sum(v is not None for v in values), known)
        self.assertEqual(sum(v is None for v in values), nulls)
        self.assertGreater(nulls, 0)  # demo data has unanswered applications

    def test_dashboard_uses_live_structured_formulas(self):
        wb = _wb()
        ws = wb["Dashboard"]
        formulas = [c.value for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        text = "\n".join(formulas)
        self.assertIn("tblApplications[", text)
        for fn in ("COUNTIFS(", "MEDIAN(", "AVERAGEIFS(", "SUMPRODUCT(", "INDEX(ResponseBands[band]"):
            self.assertIn(fn, text)
        for banned in ("LET(", "FILTER(", "XLOOKUP(", "UNIQUE(", "SORT(", "AGGREGATE("):
            self.assertNotIn(banned, text.upper().replace("_XLFN.", ""))
        self.assertEqual(len(ws._charts), 3)
        self.assertIn("kpi_response_rate", wb.defined_names)

    def test_data_dictionary_covers_every_column(self):
        wb = _wb()
        headers = [c.value for c in wb["Applications"][1]]
        ws = wb["Data dictionary"]
        cells = {c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        for header in headers:
            self.assertIn(header, cells)
            self.assertTrue(catalog.description(header), header)


@unittest.skipUnless(verify.soffice_path(), "LibreOffice (soffice) not installed")
class LibreOfficeRecalcTest(unittest.TestCase):
    def test_excel_formulas_equal_sql_metrics(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "workbook.xlsx"
            src.write_bytes(WORKBOOK)
            out = verify.recalculate(src, Path(tmp) / "recalc")
            values = verify.read_kpi_table(out)
            problems = verify.compare_dashboard(out, SETTINGS)
            wb = load_workbook(out, data_only=True)
            errors = [f"{ws.title}!{c.coordinate}" for ws in wb for row in ws.iter_rows()
                      for c in row if isinstance(c.value, str) and c.value.startswith("#")
                      and c.value.rstrip("!?0/").upper() in ("#VALUE", "#DIV", "#REF", "#NAME",
                                                             "#N/A", "#NUM", "#NULL")]
        self.assertEqual(problems, [])
        self.assertEqual(errors, [])
        with connect(SETTINGS) as con:
            kpis = Metrics(con, SETTINGS).kpis()
        self.assertEqual(values["kpis.applications"], kpis["applications"])
        self.assertAlmostEqual(values["kpis.response_rate_pct"] * 100,
                               kpis["response_rate_pct"], delta=0.051)
        self.assertEqual(values["kpis.response_band"], kpis["response_band"])
        self.assertAlmostEqual(values["kpis.median_days_to_response"],
                               kpis["median_days_to_response"])


class DaxTest(unittest.TestCase):
    def test_enough_measures(self):
        self.assertGreaterEqual(len(dax_measures()), 20)

    def test_column_references_exist_in_warehouse(self):
        columns = _warehouse_columns()
        pattern = re.compile(r"(?:'([^']+)'|\b([A-Za-z_][A-Za-z0-9_]*))\[([^\]]+)\]")
        seen = 0
        for name, expr in dax_measures().items():
            for quoted, bare, column in pattern.findall(expr):
                table = quoted or bare
                self.assertIn(table, columns, f"{name}: unknown table {table}")
                self.assertIn(column, columns[table], f"{name}: {table}[{column}] missing")
                seen += 1
        self.assertGreater(seen, 30)

    def test_measure_references_resolve(self):
        measures = dax_measures()
        bare = re.compile(r"(?<![A-Za-z0-9_'\]])\[([^\]]+)\]")
        for name, expr in measures.items():
            for ref in bare.findall(expr):
                self.assertIn(ref, measures, f"{name} references unknown measure [{ref}]")

    def test_use_relationship_for_role_playing_dates(self):
        text = "\n".join(dax_measures().values())
        self.assertIn("USERELATIONSHIP ( fct_applications[first_response_date], dim_date[date_day] )", text)
        self.assertIn("USERELATIONSHIP ( fct_applications[last_activity_date], dim_date[date_day] )", text)

    def test_blank_aware_median(self):
        expr = dax_measures()["Median Days To Response"]
        self.assertIn("NOT ISBLANK ( fct_applications[days_to_response] )", expr)

    def test_dax_file_is_a_query(self):
        text = render_dax_file()
        self.assertTrue(text.count("MEASURE fct_applications[") == len(dax_measures()))
        self.assertIn("DEFINE", text)
        self.assertIn("EVALUATE", text)
        body = re.sub(r"//.*", "", text)
        self.assertEqual(body.count("("), body.count(")"))


class TableauTest(unittest.TestCase):
    def test_enough_fields(self):
        self.assertGreaterEqual(len(tableau_calculations()), 12)

    def test_field_references_exist(self):
        calcs = tableau_calculations()
        flat = _warehouse_columns()[catalog.FLAT_VIEW]
        for name, formula in calcs.items():
            text = re.sub(r"//.*", "", formula)
            for param in re.findall(r"\[Parameters\]\.\[([^\]]+)\]", text):
                self.assertIn(param, PARAMETERS)
            text = re.sub(r"\[Parameters\]\.\[[^\]]+\]", "", text)
            for ref in re.findall(r"\[([^\]]+)\]", text):
                self.assertTrue(ref in calcs or ref in flat,
                                f"{name}: [{ref}] is neither a calculated field nor a column")

    def test_flat_view_matches_catalog(self):
        flat = _warehouse_columns()[catalog.FLAT_VIEW]
        self.assertEqual(set(catalog.column_names(catalog.FLAT_VIEW)), flat)


class BenchmarkParityTest(unittest.TestCase):
    expected = {b.label: b.lower_pct for b in bm.RESPONSE_RATE_BANDS[:-1]}

    def test_dax_bands(self):
        expr = dax_measures()["Response Band"]
        found = {label: float(v) for v, label in re.findall(r'_pct >= ([\d.]+), "([^"]+)"', expr)}
        self.assertEqual(found, self.expected)
        self.assertIn(f'"{bm.RESPONSE_RATE_BANDS[-1].label}"', expr)
        measures = dax_measures()
        self.assertAlmostEqual(float(measures["Direct Interview Benchmark %"]),
                               bm.DIRECT_INTERVIEW_RATE_PCT / 100)
        self.assertIn(f"- {bm.DIRECT_INTERVIEW_RATE_PCT}",
                      measures["Interview Rate vs Direct Benchmark (pp)"])
        conf = measures["Sample Confidence"]
        self.assertIn(f">= {bm.SAMPLE_ADEQUATE}", conf)
        self.assertIn(f">= {bm.SAMPLE_INDICATIVE}", conf)
        self.assertIn(str(Z), measures["Response Rate CI Low %"])

    def test_tableau_bands(self):
        expr = tableau_calculations()["Response Band"]
        found = {label: float(v) for v, label in re.findall(r">= ([\d.]+) THEN '([^']+)'", expr)}
        self.assertEqual(found, self.expected)
        self.assertIn(str(bm.DIRECT_INTERVIEW_RATE_PCT),
                      tableau_calculations()["Interview Rate vs Direct Benchmark (pp)"])

    def test_excel_bands_and_benchmarks(self):
        wb = _wb()
        ws = wb["Benchmarks"]

        def rows(table: str) -> list[tuple]:
            ref = ws.tables[table].ref
            return [tuple(c.value for c in row) for row in ws[ref]][1:]

        bands = {r[0]: r[1] for r in rows("ResponseBands")}
        self.assertEqual(bands, {b.label: b.lower_pct for b in bm.RESPONSE_RATE_BANDS})
        lowers = [r[1] for r in rows("ResponseBands")]
        self.assertEqual(lowers, sorted(lowers))  # approximate MATCH needs ascending order
        values = {r[0]: r[2] for r in rows("Benchmarks")}
        self.assertEqual(values["direct_interview_rate_pct"], bm.DIRECT_INTERVIEW_RATE_PCT)
        self.assertEqual(values["sample_adequate"], bm.SAMPLE_ADEQUATE)
        self.assertEqual(values["sample_indicative"], bm.SAMPLE_INDICATIVE)
        self.assertEqual(values["wilson_z"], Z)
        lo, hi = bm.FUNNEL_PER_100["Interview"]
        self.assertEqual((values["interviews_per_100_lo"], values["interviews_per_100_hi"]), (lo, hi))
        self.assertEqual(values["follow_up_days"], SETTINGS.follow_up_after_days)

    def test_model_json_benchmarks(self):
        spec = model_spec()
        self.assertEqual(spec["benchmarks"]["response_rate_bands_pct"],
                         {b.label: b.lower_pct for b in bm.RESPONSE_RATE_BANDS})


class ModelSpecTest(unittest.TestCase):
    def test_tables_and_columns_match_warehouse(self):
        columns = _warehouse_columns()
        for table in model_spec()["tables"]:
            names = {c["name"] for c in table["columns"]}
            self.assertEqual(names, columns[table["name"]], table["name"])

    def test_relationships(self):
        spec = model_spec()
        columns = _warehouse_columns()
        rels = spec["relationships"]
        for rel in rels:
            self.assertIn(rel["fromColumn"], columns[rel["fromTable"]])
            self.assertIn(rel["toColumn"], columns[rel["toTable"]])
            self.assertEqual(rel["cardinality"], "manyToOne")
        date_rels = [r for r in rels if r["toTable"] == "dim_date"]
        self.assertEqual(sum(r["isActive"] for r in date_rels), 1)
        inactive = {r["fromColumn"] for r in date_rels if not r["isActive"]}
        self.assertEqual(inactive, {"first_response_date", "last_activity_date"})
        date_table = next(t for t in spec["tables"] if t["name"] == "dim_date")
        self.assertTrue(date_table["isMarkedDateTable"])

    def test_theme(self):
        theme = powerbi_theme()
        docs_copy = json.loads((ROOT / "docs" / "bi" / "powerbi-theme.json").read_text(encoding="utf-8"))
        self.assertEqual(docs_copy, theme, "regenerate docs/bi/powerbi-theme.json")
        self.assertEqual(theme["background"], "#F7F5F0")
        self.assertEqual(theme["dataColors"][0], theme["tableAccent"])
        self.assertGreaterEqual(len(theme["dataColors"]), 8)
        self.assertTrue(all(re.fullmatch(r"#[0-9A-F]{6}", c) for c in theme["dataColors"]))


class ExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.out = Path(tempfile.mkdtemp(prefix="bi-export-"))
        cls.written = export_all(SETTINGS, cls.out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.out, ignore_errors=True)

    def test_files_written(self):
        for key in ("workbook", "job_applications.csv", "powerbi/measures.dax",
                    "powerbi/theme.json", "powerbi/model.json", "tableau/calculations.md",
                    "README.md"):
            self.assertIn(key, self.written)
            self.assertTrue(self.written[key].exists(), key)
            self.assertGreater(self.written[key].stat().st_size, 0)
        for table in catalog.TABLES:
            self.assertTrue((self.out / "star_schema" / f"{table}.csv").exists(), table)
        json.loads((self.out / "powerbi" / "theme.json").read_text(encoding="utf-8"))
        model = json.loads((self.out / "powerbi" / "model.json").read_text(encoding="utf-8"))
        self.assertGreaterEqual(len(model["measures"]), 20)
        load_workbook(self.written["workbook"])

    def test_csv_keeps_unknown_reply_times_blank(self):
        lines = (self.out / "job_applications.csv").read_text(encoding="utf-8").splitlines()
        header = lines[0].split(",")
        self.assertIn("days_to_response", header)
        import csv
        rows = list(csv.DictReader(lines))
        blank = sum(1 for r in rows if r["days_to_response"] == "")
        with connect(SETTINGS) as con:
            nulls = con.execute("SELECT COUNT(*) - COUNT(days_to_response) "
                                "FROM rpt_application_analysis").fetchone()[0]
        self.assertEqual(blank, nulls)

    def test_no_insurance_measures_anywhere(self):
        texts = {k: p.read_text(encoding="utf-8") for k, p in self.written.items()
                 if p.suffix in (".dax", ".json", ".md")}
        texts["dax"] = json.dumps(dax_measures())
        texts["tableau"] = json.dumps(tableau_calculations())
        for doc in (ROOT / "docs" / "bi").glob("*"):
            if doc.is_file() and doc.suffix in (".md", ".json"):
                texts[str(doc)] = doc.read_text(encoding="utf-8")
        wb = load_workbook(self.written["workbook"])
        for sheet in ("Cover", "Dashboard", "Data dictionary", "Benchmarks"):
            texts[sheet] = "\n".join(str(c.value) for row in wb[sheet].iter_rows() for c in row
                                     if c.value is not None)
        for name, text in texts.items():
            self.assertIsNone(INSURANCE_TERMS.search(text), f"insurance term in {name}")


if __name__ == "__main__":
    unittest.main()
